"""Tests for text extraction and input safeguards."""

from __future__ import annotations

import io

import pytest

from courselens import config
from courselens.errors import (
    EmptyInputError,
    ScannedDocumentError,
    UnsupportedFileTypeError,
)
from courselens.extraction import (
    extract_from_file,
    extract_from_text,
)


def test_extract_from_text_basic():
    result = extract_from_text("  Hello CourseLens  ")
    assert result.text == "Hello CourseLens"
    assert result.source == "text"
    assert result.truncated is False


def test_extract_from_text_empty_raises():
    with pytest.raises(EmptyInputError):
        extract_from_text("   \n  \t ")


def test_extract_from_text_normalizes_line_endings():
    result = extract_from_text("line1\r\nline2\rline3")
    assert result.text == "line1\nline2\nline3"


def test_truncation_applies_size_cap():
    big = "a" * (config.MAX_INPUT_CHARS + 500)
    result = extract_from_text(big)
    assert result.truncated is True
    assert len(result.text) == config.MAX_INPUT_CHARS
    assert result.original_chars == config.MAX_INPUT_CHARS + 500


def test_extract_txt_file():
    data = "# Notes\nSome markdown content".encode("utf-8")
    result = extract_from_file(data, "notes.md")
    assert "markdown content" in result.text
    assert result.source == "txt"


def test_extract_txt_utf8_bom():
    # encode("utf-8-sig") prepends the BOM bytes; the extractor should strip
    # them so the text starts cleanly.
    data = "with bom".encode("utf-8-sig")
    result = extract_from_file(data, "x.txt")
    assert result.text == "with bom"


def test_unsupported_file_type_raises():
    with pytest.raises(UnsupportedFileTypeError):
        extract_from_file(b"data", "sheet.csv")


def test_scanned_pdf_detection(monkeypatch):
    # Simulate a PDF that parses but yields almost no text.
    import courselens.extraction as extraction

    monkeypatch.setattr(extraction, "_extract_pdf", lambda data: "  ")
    with pytest.raises(ScannedDocumentError):
        extract_from_file(b"%PDF-1.4 fake", "scanned.pdf")


def test_pdf_roundtrip_real_library():
    # Build a tiny real PDF with pypdf so we exercise the actual extractor.
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    # pypdf can't easily inject text; write a blank page and confirm the
    # scanned-document guard triggers (blank page -> no text).
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ScannedDocumentError):
        extract_from_file(buf.getvalue(), "blank.pdf")


def test_docx_roundtrip_real_library():
    pytest.importorskip("docx")
    import docx

    document = docx.Document()
    document.add_paragraph("CourseLens assignment: build a study plan.")
    document.add_paragraph("Due next Friday.")
    buf = io.BytesIO()
    document.save(buf)

    result = extract_from_file(buf.getvalue(), "assignment.docx")
    assert "study plan" in result.text
    assert "Due next Friday" in result.text
    assert result.source == "docx"
